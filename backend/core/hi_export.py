"""Health-index export: score a chosen set of assets on chosen criteria, to Excel.

A fleet export can take minutes (every asset is a live calculation against a
slow database), so it runs as a background job: `start()` returns a job id, the
UI polls `status()`, and `workbook_path()` hands over the finished file.

Each asset is scored by :func:`hi_engine.compute` exactly as the manual
calculation screen does it - the chosen criteria, weights re-normalised over
whatever has data, the optional DGA correction factor - and the oil-filtration
rule for DGA level and trend (core/dga_filter) applies automatically.
"""
from __future__ import annotations

import logging
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any

import db
from config import BASE_DIR, Config
from core import assets as assets_mod
from core import dga_filter, hi_engine, hierarchy, scoring

log = logging.getLogger(__name__)

EXPORT_DIR = BASE_DIR / "snapshots" / "exports"
WORKERS = 4                 # parallel calculations; the database is the bottleneck
KEEP_JOBS = 20              # finished jobs (and their files) kept for download

_jobs: dict[str, dict[str, Any]] = {}
_lock = threading.Lock()


# --------------------------------------------------------------------------
# What can be exported
# --------------------------------------------------------------------------
CATEGORY_TYPE = {"TR": "TR", "AET": "TR", "OLTC": "TR", "CTVT": "CTVT", "CB": "CB",
                 "ESDS": "ESDS", "SA": "SA", "CBank": "CBank", "BBank": "BBank"}


def options() -> dict[str, Any]:
    """Categories with their asset counts, and each scoring type's criteria."""
    counts: dict[str, int] = {}
    for a in assets_mod.all_assets():
        counts[a.category] = counts.get(a.category, 0) + 1
    weights = scoring.weight_sets()
    criteria = {
        t: [{"code": c, "label": scoring.COMPONENT_LABEL.get(c, c),
             "manual": c in scoring.MANUAL_COMPONENTS.get(t, set()),
             "configuredWeight": weights.get(t, {}).get(c, 0.0)}
            for c in order]
        for t, order in scoring.COMPONENT_ORDER.items()
    }
    return {
        "categories": [{"key": k, "assetType": t, "count": counts.get(k, 0)}
                       for k, t in CATEGORY_TYPE.items() if counts.get(k)],
        "criteria": criteria,
        "dgaCorrectionDefaults": {str(k): v for k, v in hi_engine.DGA_CORRECTION_DEFAULTS.items()},
        "filterWindow": {"enabled": Config.DGA_FILTER_WINDOW,
                         "minSamplesAfter": Config.DGA_MIN_SAMPLES_AFTER_FILTER},
    }


def resolve_assets(category: str, selected: list[str] | None, node: str = "",
                   excluded: set[str] | None = None) -> list[str]:
    """The asset numbers to export: the chosen ones, or every asset in the category."""
    if selected:
        known = {a.assetNumber.upper(): a.assetNumber for a in assets_mod.all_assets()}
        out, seen = [], set()
        for s in selected:
            n = known.get((s or "").strip().upper())
            if n and n not in seen:
                out.append(n)
                seen.add(n)
        return out
    excluded = excluded or set()
    items = [a for a in assets_mod.all_assets() if a.category == category
             and a.assetNumber not in excluded
             and (not node or a.assetNumber == node or a.assetNumber.startswith(node + "/"))]
    return sorted(a.assetNumber for a in items)


# --------------------------------------------------------------------------
# Jobs
# --------------------------------------------------------------------------
def start(category: str, assets: list[str], criteria: list[str] | None,
          dga_factors: dict[int, float] | None, date_from: str, date_to: str,
          user: str = "engineer", scope_label: str | None = None) -> dict[str, Any]:
    if category not in CATEGORY_TYPE:
        raise ValueError(f"unknown asset category '{category}'")
    if not assets:
        raise ValueError("no assets to export")
    asset_type = CATEGORY_TYPE[category]
    order = scoring.COMPONENT_ORDER.get(asset_type, [])
    if criteria is not None:
        bad = [c for c in criteria if c not in order]
        if bad:
            raise ValueError(f"criteria not defined for {asset_type}: {', '.join(bad)}")
        if not criteria:
            raise ValueError("select at least one criterion")
    if dga_factors is not None and asset_type != "TR":
        dga_factors = None                      # the DGA factor is a transformer-only idea

    job_id = uuid.uuid4().hex[:12]
    job = {
        "id": job_id, "status": "running", "category": category, "assetType": asset_type,
        "total": len(assets), "done": 0, "failed": 0, "startedAt": datetime.now().isoformat(timespec="seconds"),
        "finishedAt": None, "error": None, "file": None, "fileName": None, "user": user,
        "settings": {"criteria": criteria, "dgaFactors": dga_factors, "from": date_from, "to": date_to,
                     "scope": scope_label, "assetCount": len(assets)},
    }
    with _lock:
        _jobs[job_id] = job
        _trim()
    threading.Thread(target=_run, args=(job, assets, criteria, dga_factors, date_from, date_to),
                     daemon=True, name=f"hi-export-{job_id}").start()
    return public(job)


def _trim() -> None:
    finished = sorted((j for j in _jobs.values() if j["status"] != "running"),
                      key=lambda j: j["startedAt"])
    stale = finished[:-KEEP_JOBS] if len(finished) > KEEP_JOBS else []
    for j in stale:
        _jobs.pop(j["id"], None)
        if j.get("file"):
            Path(j["file"]).unlink(missing_ok=True)


def public(job: dict[str, Any]) -> dict[str, Any]:
    return {k: v for k, v in job.items() if k != "file"} | {"ready": job["status"] == "complete"}


def status(job_id: str) -> dict[str, Any] | None:
    with _lock:
        job = _jobs.get(job_id)
        return public(job) if job else None


def workbook_path(job_id: str) -> tuple[Path, str] | None:
    with _lock:
        job = _jobs.get(job_id)
    if not job or job["status"] != "complete" or not job.get("file"):
        return None
    return Path(job["file"]), job["fileName"]


def _score_one(asset: str, criteria, dga_factors, date_from, date_to) -> dict[str, Any]:
    last: Exception | None = None
    for _ in range(3):                           # the live database times out now and then
        try:
            return hi_engine.compute(asset, date_from, date_to, selected=criteria,
                                     dga_correction=dga_factors)
        except Exception as exc:                 # noqa: BLE001 - one asset must not end the job
            last = exc
            time.sleep(1)
    return {"asset": asset, "error": str(last)}


def _run(job, assets, criteria, dga_factors, date_from, date_to) -> None:
    results: dict[str, dict[str, Any]] = {}
    try:
        with ThreadPoolExecutor(max_workers=WORKERS) as pool:
            futures = {pool.submit(_score_one, a, criteria, dga_factors, date_from, date_to): a
                       for a in assets}
            for fut in as_completed(futures):
                a = futures[fut]
                res = fut.result()
                results[a] = res
                with _lock:
                    job["done"] += 1
                    if res.get("error"):
                        job["failed"] += 1
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
        name = f"HI_Export_{job['category']}_{stamp}.xlsx"
        path = EXPORT_DIR / f"{job['id']}.xlsx"
        build_workbook(path, job, [results[a] for a in assets])
        with _lock:
            job.update(status="complete", file=str(path), fileName=name,
                       finishedAt=datetime.now().isoformat(timespec="seconds"))
    except Exception as exc:                     # noqa: BLE001
        log.exception("HI export %s failed", job["id"])
        with _lock:
            job.update(status="failed", error=str(exc),
                       finishedAt=datetime.now().isoformat(timespec="seconds"))


# --------------------------------------------------------------------------
# Workbook
# --------------------------------------------------------------------------
def _names(assets: list[str]) -> dict[str, str]:
    """CMMS names (`ast_mst_asset_shortdesc`) for the assets and their sites.

    Read straight from the CMMS for just these asset numbers - this runs in a
    background job, so the long timeout is acceptable - and fall back to the
    navigator's cached tree when the CMMS cannot be reached.
    """
    wanted = sorted({a for a in assets if a} | {a.split("/")[0] for a in assets if a})
    out: dict[str, str] = {}
    try:
        for i in range(0, len(wanted), 400):
            batch = wanted[i:i + 400]
            rows = db.tomms_query_long(
                "SELECT LTRIM(RTRIM(ast_mst_asset_no)) AS a, LTRIM(RTRIM(ast_mst_asset_shortdesc)) AS d "
                f"FROM ast_mst WHERE site_cd = ? AND ast_mst_asset_no IN ({', '.join('?' * len(batch))})",
                [Config.TOMMS_SITE_CD, *batch], timeout=180)
            out.update({r["a"]: r["d"] for r in rows if r.get("a") and r.get("d")})
    except Exception as exc:                     # noqa: BLE001 - names are a nicety
        log.warning("export: CMMS names unavailable, using the navigator: %s", exc)
    for a in wanted:
        if a in out:
            continue
        try:
            node = hierarchy.node(a)
        except Exception:                        # noqa: BLE001
            break
        if node and node.get("named"):
            out[a] = node["label"]
    return out


def _fmt_window(comp: dict[str, Any]) -> str:
    w = (comp.get("detail") or {}).get("filterWindow")
    return w["reason"] if w else ""


def build_workbook(path: Path, job: dict[str, Any], results: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as col

    F = "Arial"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    grey = PatternFill("solid", fgColor="F2F2F2")

    def score_fill(s):
        if s is None:
            return grey
        return PatternFill("solid", fgColor="E6F6F0" if s >= .75 else "E8F1FC" if s >= .5
                           else "FDF6E3" if s >= .25 else "FCEAEA")

    def header(ws, row, labels):
        for i, text in enumerate(labels, 1):
            c = ws.cell(row, i, text)
            c.font = Font(name=F, bold=True, color="FFFFFF")
            c.fill = head_fill
            c.border = border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 34

    def put(ws, row, column, value, fmt=None, bold=False, fill=None, left=False, color=None):
        c = ws.cell(row, column, value)
        c.font = Font(name=F, bold=bold, color=color)
        c.border = border
        c.alignment = Alignment(horizontal="left" if left else "center", vertical="center")
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        return c

    asset_type = job["assetType"]
    order = scoring.COMPONENT_ORDER.get(asset_type, [])
    criteria = job["settings"]["criteria"] or order
    criteria = [c for c in order if c in criteria]
    labels = {c: scoring.COMPONENT_LABEL.get(c, c) for c in criteria}
    dga_on = job["settings"]["dgaFactors"] is not None
    names = _names([r.get("asset") for r in results if r.get("asset")])

    wb = Workbook()
    ws = wb.active
    ws.title = "HI Summary"
    ws["A1"] = f"Health Index Export – {job['category']}"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    crit_text = "all criteria" if not job["settings"]["criteria"] else ", ".join(labels.values())
    ws["A2"] = (f"Calculated {datetime.now():%Y-%m-%d %H:%M} · {len(results)} assets · {crit_text}"
                f" · DGA correction {'ON' if dga_on else 'OFF'}"
                f" · oil-filtration rule {'ON' if Config.DGA_FILTER_WINDOW else 'OFF'}")
    ws["A2"].font = Font(name=F, italic=True, size=9, color="595959")
    ws["A3"] = ("Criterion cells are scores (0–1, 1 = healthy). Grey 'no data' criteria are left out "
                "and the weights re-normalise over the rest. See the Settings sheet for the method.")
    ws["A3"].font = Font(name=F, italic=True, size=9, color="595959")

    cols = ["No", "Asset Number", "Name", "Site", "Type"] + [f"{labels[c]}" for c in criteria]
    if dga_on:
        cols += ["HI before DGA correction", "DGA status", "Correction factor"]
    cols += ["Health Index", "Condition", "Criteria used", "No data for", "DGA data window", "Note"]
    H = 5
    header(ws, H, cols)
    for i, r in enumerate(results, 1):
        row = H + i
        a = r.get("asset", "")
        comps = {c["code"]: c for c in r.get("components", [])}
        put(ws, row, 1, i)
        put(ws, row, 2, a, left=True, bold=True)
        put(ws, row, 3, names.get(a, ""), left=True)
        site = a.split("/")[0] if a else ""
        put(ws, row, 4, f"{site} {names.get(site, '')}".strip(), left=True)
        put(ws, row, 5, r.get("assetTypeLabel", ""), left=True)
        k = 6
        for c in criteria:
            comp = comps.get(c, {})
            s = comp.get("score") if comp.get("available") else None
            put(ws, row, k, s if s is not None else "no data", "0.00", fill=score_fill(s),
                color=None if s is not None else "808080")
            k += 1
        corr = r.get("dgaCorrection") or {}
        if dga_on:
            put(ws, row, k, r.get("healthIndexUncorrected"), "0.00")
            put(ws, row, k + 1, corr.get("statusLabel") or "not classified", left=True)
            put(ws, row, k + 2, corr.get("factor") if corr.get("applied") else None, "0.00")
            k += 3
        band = r.get("band") or {}
        put(ws, row, k, r.get("healthIndex"), "0.00", bold=True)
        put(ws, row, k + 1, band.get("label") or "", bold=True,
            color=(band.get("color") or "#000000").lstrip("#"))
        used = sum(1 for c in criteria if comps.get(c, {}).get("available"))
        put(ws, row, k + 2, f"{used} / {len(criteria)}")
        put(ws, row, k + 3, ", ".join(labels[c] for c in criteria
                                      if not comps.get(c, {}).get("available")) or "—", left=True)
        windows = {_fmt_window(comps.get(c, {})) for c in ("DGAL", "DGAT") if c in criteria}
        put(ws, row, k + 4, "; ".join(w for w in windows if w) or ("full history" if windows else ""),
            left=True)
        note = r.get("error") or r.get("message") or ""
        if corr.get("message") and dga_on:
            note = f"{note} {corr['message']}".strip()
        put(ws, row, k + 5, note, left=True)
    widths = [5, 27, 32, 22, 22] + [11] * len(criteria) + ([12, 14, 11] if dga_on else []) \
        + [11, 12, 10, 36, 46, 36]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[col(i)].width = w
    ws.freeze_panes = ws.cell(H + 1, 3)
    ws.auto_filter.ref = f"A{H}:{col(len(cols))}{H + len(results)}"

    # --- per-criterion detail
    wd = wb.create_sheet("Criteria Detail")
    header(wd, 1, ["Asset Number", "Criterion", "Measured value", "Unit", "Test date", "Score (0–1)",
                   "Normalised weight %", "Sub HI", "Source", "How the value was chosen"])
    row = 2
    band_fill = PatternFill("solid", fgColor="DDEBF7")
    for r in results:
        comps = {c["code"]: c for c in r.get("components", [])}
        for c in criteria:
            comp = comps.get(c, {})
            av = bool(comp.get("available"))
            put(wd, row, 1, r.get("asset"), left=True, bold=True)
            put(wd, row, 2, labels[c], left=True)
            put(wd, row, 3, comp.get("value") if av else "no data", "0.000" if av else None,
                color=None if av else "808080")
            put(wd, row, 4, comp.get("unit") or "")
            put(wd, row, 5, str(comp.get("date"))[:10] if av and comp.get("date") else "—")
            put(wd, row, 6, comp.get("score") if av else "—", "0.00",
                fill=score_fill(comp.get("score") if av else None))
            put(wd, row, 7, comp.get("weight") if av else "—", "0.00")
            put(wd, row, 8, comp.get("contribution") if av else "—", "0.00")
            put(wd, row, 9, comp.get("source") or "", left=True)
            put(wd, row, 10, comp.get("rule") or "", left=True)
            row += 1
        for column in range(1, 11):
            put(wd, row, column, None, fill=band_fill)
        put(wd, row, 1, r.get("asset"), left=True, bold=True, fill=band_fill)
        put(wd, row, 2, "Health Index", left=True, bold=True, fill=band_fill)
        put(wd, row, 8, r.get("healthIndex"), "0.00", bold=True, fill=band_fill)
        put(wd, row, 9, (r.get("band") or {}).get("label") or "", left=True, bold=True, fill=band_fill)
        if dga_on and (r.get("dgaCorrection") or {}).get("applied"):
            put(wd, row, 10, r["dgaCorrection"]["message"], left=True, fill=band_fill)
        row += 1
    for i, w in enumerate([27, 24, 14, 10, 12, 11, 13, 9, 34, 70], 1):
        wd.column_dimensions[col(i)].width = w
    wd.freeze_panes = "B2"

    # --- settings / method
    st = wb.create_sheet("Settings")
    st["A1"] = "Export settings and method"
    st["A1"].font = Font(name=F, bold=True, size=13)
    f = job["settings"]["dgaFactors"]
    lines = [
        ("Exported", f"{datetime.now():%Y-%m-%d %H:%M} by {job.get('user') or 'engineer'}"),
        ("Asset category", f"{job['category']} (scored as {asset_type})"),
        ("Scope", job["settings"].get("scope") or "all sites"),
        ("Assets", f"{len(results)} ({job.get('failed', 0)} failed)"),
        ("Criteria", crit_text),
        ("Test date range", f"{job['settings']['from']} to {job['settings']['to']}"),
        ("Weights", "Configured weights re-normalised over the selected criteria that have data: "
                    "w / Σw(available) × Σw(all)."),
        ("DGA correction factor", "OFF" if f is None else
         "ON – final HI × (1 − factor) by IEEE C57.104-2019 status: " +
         ", ".join(f"Status {k} = {v:g}" for k, v in sorted(f.items()))),
        ("Oil-filtration rule", "OFF" if not Config.DGA_FILTER_WINDOW else
         f"ON – DGA level and trend use the samples after the last filter point when there are at "
         f"least {Config.DGA_MIN_SAMPLES_AFTER_FILTER}; otherwise the data before it (repeated for "
         "earlier points). Filter points: detected (not rejected) plus engineer-marked."),
        ("Filter-point detection", dga_filter.RULE),
    ]
    for i, (k, v) in enumerate(lines, start=3):
        st.cell(i, 1, k).font = Font(name=F, bold=True)
        c = st.cell(i, 2, v)
        c.font = Font(name=F)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    st.column_dimensions["A"].width = 24
    st.column_dimensions["B"].width = 110
    wb.save(path)
