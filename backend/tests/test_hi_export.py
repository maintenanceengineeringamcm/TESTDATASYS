"""Health-index export: job validation and the workbook (unit, no database)."""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from core import hi_export


def comp(code, score, value=1.0, available=True, rule="", detail=None):
    return {"code": code, "label": code, "available": available, "score": score if available else None,
            "value": value if available else None, "unit": "", "date": "2026-01-01", "weight": 50.0,
            "contribution": (score or 0) * 50.0, "source": "T", "rule": rule, "detail": detail or {}}


RESULT = {
    "asset": "G001/PE/1/02/MT01_X/01", "assetTypeLabel": "Power Transformer", "healthIndex": 21.0,
    "healthIndexUncorrected": 70.0, "band": {"label": "Very Poor", "color": "#D64545"},
    "dgaCorrection": {"applied": True, "statusLabel": "Status 3", "factor": 0.7,
                      "message": "Status 3: index reduced by 70%"},
    "components": [
        comp("AGE", 0.5, 19.0),
        comp("DGAL", 0.9, 0.8, rule="weighted index · 4 samples since the filter point of 2025-01-06",
             detail={"filterWindow": {"reason": "4 samples since the filter point of 2025-01-06"}}),
        comp("FURAN", None, available=False),
    ],
}


def test_workbook_layout(tmp_path, monkeypatch):
    monkeypatch.setattr(hi_export, "_names", lambda assets: {"G001/PE/1/02/MT01_X/01": "TR 1",
                                                             "G001": "Ambalangoda GS"})
    job = {"id": "x", "category": "TR", "assetType": "TR", "failed": 0, "user": "eng",
           "settings": {"criteria": ["AGE", "FURAN", "DGAL"], "dgaFactors": {1: 0.0, 2: 0.3, 3: 0.7},
                        "from": "1900-01-01", "to": "2999-12-31", "scope": None}}
    path = tmp_path / "x.xlsx"
    hi_export.build_workbook(path, job, [RESULT])
    wb = load_workbook(path)
    assert wb.sheetnames == ["HI Summary", "Criteria Detail", "Settings"]
    ws = wb["HI Summary"]
    head = [c.value for c in ws[5]]
    # criteria in engine order (AGE, FURAN, DGAL), then the DGA correction columns
    assert head[5:8] == ["Age", "Furanic Level / DP", "DGA Gas Levels"]
    assert head[8:15] == ["DGA status", "DGA status reason", "Correction factor",
                          "HI without DGA correction", "Condition without correction",
                          "HI with DGA correction", "Condition with correction"]
    row = {h: c.value for h, c in zip(head, ws[6])}
    assert row["Name"] == "TR 1" and row["Site"] == "G001 Ambalangoda GS"
    assert row["Furanic Level / DP"] == "no data" and row["Age"] == 0.5
    assert row["HI with DGA correction"] == 21.0 and row["HI without DGA correction"] == 70.0
    assert row["Correction factor"] == 0.7 and row["Criteria used"] == "2 / 3"
    assert "filter point of 2025-01-06" in row["DGA data window"]
    detail = wb["Criteria Detail"]
    assert detail.cell(2, 1).value == "G001/PE/1/02/MT01_X/01" and detail.cell(5, 2).value == "HI with DGA correction"


def test_start_validation():
    with pytest.raises(ValueError):
        hi_export.start("NOPE", ["a"], None, None, "1900-01-01", "2999-12-31")
    with pytest.raises(ValueError):
        hi_export.start("TR", [], None, None, "1900-01-01", "2999-12-31")
    with pytest.raises(ValueError):
        hi_export.start("CB", ["a"], ["DGAL"], None, "1900-01-01", "2999-12-31")
    with pytest.raises(ValueError):
        hi_export.start("TR", ["a"], [], None, "1900-01-01", "2999-12-31")
